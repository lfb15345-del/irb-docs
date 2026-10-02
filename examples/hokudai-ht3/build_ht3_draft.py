# -*- coding: utf-8 -*-
"""北大病院 HT3「観察研究用 情報公開用文書 雛形（2025年4月版）」から下書きを作る。

赤字（記載上の注意）と青字（例文）を消し、見出しは残し、
本文が空になった見出しの下に「（本文を記載）」を1行置く。

    python build_ht3_draft.py --template 様式/HT3_情報公開用文書ひな形_202504.docx   # ①②の両ページを残す
    python build_ht3_draft.py --template … --variant multi      # ②だけ
    python build_ht3_draft.py --template … --variant single     # ①だけ
    python build_ht3_draft.py --template … --out 別名.docx --report 記録.json

HT3 は HT1/HT2 と作りが違うので、docx_tools の purge_guidance / purge_examples / Inserter は使わない。
  - 見出しは「[研究課題名]」「○対象となる患者さん」の形。docx_tools.HEADING_RX（1．/（1）/【】）では拾えない。
  - 1行に黒字・赤字・青字が混ざるので、段落単位ではなくラン（同じ書式の文字のかたまり）単位で処理する。
      赤字 → 消す
      青字 → 「／研究用に保管された検体」「検体・」のような選択肢は消す。
              それ以外（○○科、011-・・・、20○○年 など）は【　　】に置き換える（黒字の定型文・ラベルは残す）
      段落が赤字・青字だけ → 段落ごと消す
      黒字の「＊20○○年○○月○○日までのカルテ情報を収集します。」 → 例文として段落ごと消す
      「＊上記の研究に…ご連絡ください。」は拒否の機会を知らせる定型文なので残す（先頭の＊だけ外す）
  - 表（共同研究機関・既存試料情報の提供のみを行う機関）は見出し行を黒にし、「例）」の行を消し、空の行に（本文を記載）。
  - ヘッダーの「20○○年○月○日（第〇版）」は【　　】に。冒頭の注意ページを消すと1ページ目に
    「先頭ページのみ別指定」の空ヘッダーが出るので、その指定を外す。
  - ①と②のページに同じ見出しがあるので、見出しは「最初の一致」ではなく全部の一致を処理する。
  - 位置は見出しの文字で決める（段落の番号は使わない）。

ひな形は読むだけ。出力の既定はコマンドを実行したフォルダ（--out で指定できる）。
2026-09-17 にスキルの試験で書かれた下書き用スクリプト。docx_tools v2 の mode="inline" 等はこの処理を一般化したもの。
ひな形と同じフォルダには書かない（ひな形や提出物を上書きしないため）。
"""
import argparse
import copy
import io
import json
import re
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))   # この skill の scripts/

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from docx_tools import (BLUE, GREEN_ETC, RED, blocks, clear_marks, cut_before, cut_from,  # noqa: E402
                        distinct_cells, drop_leading_blanks, find, insert_after, rcolor, remove)

TPL = None   # --template で渡す（北大の HT3 雛形 2025年4月版）
HERE = Path(__file__).resolve().parent

PLACEHOLDER = "（本文を記載）"
SLOT = "【　　】"
BODY_PT = 12                      # HT3 の本文は 12pt
GUIDE_BLUE = BLUE | GREEN_ETC

VARIANT_1 = r"^①当院単独で研究を実施$"
VARIANT_2 = r"^②当院主導で多機関共同研究を実施$"
TITLE_RX = re.compile(r"^臨床研究に関する情報$")
H1_RX = re.compile(r"^\[[^\]]+\]")            # [研究課題名] など
H2_RX = re.compile(r"^[○〇](?![○〇])\S")      # ○対象となる患者さん など（「○○科に、…」は除く）
CHOICE_WORDS = ("検体・",)                    # 「検体・情報」の「検体・」＝検体を使うときだけ残す選択肢
KEEP_STAR_RX = re.compile(r"^＊上記の研究に.*ご了解いただけない場合")
BLACK_EXAMPLE = ("○○", "〇〇", "◯◯", "（例", "例）")
PARTICLE_RX = re.compile(r"[ぁ-ん、]{1,2}")   # 青字と青字の間の「に、」は例文の一部として扱う
TRACE = ("○", "〇", "◯", "＊", "例）", "（例", "・・・", "記載してください")


# ------------------------------------------------------------------ 判定
def cls_of(run):
    c = rcolor(run)
    if c in RED:
        return "red"
    if c in GUIDE_BLUE:
        return "blue"
    return "black"


def heading_kind(text):
    t = text.strip()
    if TITLE_RX.match(t):
        return "title"
    if H1_RX.match(t):
        return "h1"
    if H2_RX.match(t) and len(t) < 40:
        return "h2"
    return None


def has_page_break(p):
    return any(br.get(qn("w:type")) == "page" for br in p._p.iter(qn("w:br")))


def segments(p):
    """[色, [ラン…]] の並び。同じ色の連続はまとめ、青字に挟まれた助詞・読点は青字側に含める。"""
    segs = []
    for r in p.runs:
        if not r.text:
            continue
        c = cls_of(r)
        if segs and segs[-1][0] == c:
            segs[-1][1].append(r)
        else:
            segs.append([c, [r]])
    merged, i = [], 0
    while i < len(segs):
        c, rs = segs[i]
        if (c == "black" and merged and merged[-1][0] == "blue" and i + 1 < len(segs)
                and segs[i + 1][0] == "blue" and PARTICLE_RX.fullmatch("".join(r.text for r in rs))):
            merged[-1][1].extend(rs + segs[i + 1][1])
            i += 2
            continue
        merged.append([c, list(rs)])
        i += 1
    return merged


def seg_text(seg):
    return "".join(r.text for r in seg[1])


# ------------------------------------------------------------------ 書き換え
def clear_mark_color(p):
    """段落記号の色（赤・青）を落とす。残すと、その行の末尾で改行したとき青字で入力される。"""
    ppr = p._p.pPr
    rpr = None if ppr is None else ppr.find(qn("w:rPr"))
    if rpr is None:
        return
    col = rpr.find(qn("w:color"))
    if col is not None and (col.get(qn("w:val")) or "").upper() in RED | GUIDE_BLUE:
        rpr.remove(col)
    hl = rpr.find(qn("w:highlight"))
    if hl is not None:
        rpr.remove(hl)


def rstrip_runs(p):
    for r in reversed(p.runs):
        if r.text.strip():
            r.text = r.text.rstrip()
            return
        if r.text:
            remove(r)


def strip_leading_star(p):
    for r in p.runs:
        if r.text:
            if r.text.startswith(("＊", "*")):
                r.text = r.text[1:]
            return


def apply_inline(p, whole_colored_heading=False):
    """赤字を消し、青字は選択肢なら消す・それ以外は【　　】にする。行った変更を返す。"""
    log = []
    for c, rs in segments(p):
        text = "".join(r.text for r in rs)
        if c == "black":
            continue
        if c == "red":
            for r in rs:
                remove(r)
            log.append(["赤字を削除", text])
        elif whole_colored_heading or not text.strip():
            for r in rs:
                clear_marks(r)
            if whole_colored_heading:
                log.append(["見出しの色を黒に", text])
        elif text.startswith("／") or text in CHOICE_WORDS:
            for r in rs:
                remove(r)
            log.append(["青字の選択肢を削除", text])
        else:
            rs[0].text = SLOT
            clear_marks(rs[0])
            for r in rs[1:]:
                remove(r)
            log.append(["青字を【　　】に", text])
    clear_mark_color(p)
    return log


# ------------------------------------------------------------------ 手順
def count_blocks(bl):
    return {"blocks": len(bl),
            "nonempty": sum(1 for b in bl if isinstance(b, Table) or b.text.strip()),
            "texts": [b.text.strip()[:40] for b in bl if isinstance(b, Paragraph) and b.text.strip()]}


def drop_trailing_blanks(doc):
    while True:
        bl = blocks(doc)
        if not bl or isinstance(bl[-1], Table) or bl[-1].text.strip():
            return
        remove(bl[-1])


def cut_variant(doc, variant, stats):
    bl = blocks(doc)
    i1, i2 = find(bl, VARIANT_1), find(bl, VARIANT_2)
    stats["cover"] = count_blocks(bl[:i1])
    if variant == "multi":
        stats["other_variant"] = count_blocks(bl[i1:i2])
        cut_before(doc, VARIANT_2)            # 冒頭の注意ページ＋①のページ
    else:
        cut_before(doc, VARIANT_1)            # 冒頭の注意ページ
        if variant == "single":
            stats["other_variant"] = count_blocks(bl[i2:])
            cut_from(doc, VARIANT_2)          # ②のページ
            drop_trailing_blanks(doc)         # ②の前の改ページ段落（残すと末尾が白紙）
    drop_leading_blanks(doc)


def purge(doc, stats):
    """本文の段落をラン単位で片付ける。見出し（w:p）と、見出しごとの書式見本を返す。"""
    page = None
    current = None
    style_ref = {}              # 見出しの w:p → 消した本文の段落の写し（字下げの見本）
    headings = {}               # 見出し/表の要素 → (ページ, 種類)  ※要素を保持して同一性を保つ
    for b in blocks(doc):
        if isinstance(b, Table):
            headings[b._tbl] = (page, "table")
            continue
        t = b.text.strip()
        if re.match(VARIANT_1, t):
            page = "①"
        elif re.match(VARIANT_2, t):
            page = "②"
        if not t:
            clear_mark_color(b)
            continue
        segs = [s for s in segments(b) if seg_text(s).strip()]
        kinds = {c for c, _ in segs}
        kind = heading_kind(t)
        if kind:
            before = b.text
            log = apply_inline(b, whole_colored_heading="black" not in kinds)
            if log:
                rstrip_runs(b)
            if b.text != before:
                stats["partial"].append({"page": page, "before": before.strip(), "after": b.text.strip(), "changes": log})
            current = b._p
            headings[b._p] = (page, kind)
            continue

        reason = None
        before = b.text
        star_log = []
        if KEEP_STAR_RX.match(t):
            strip_leading_star(b)
            star_log = [["先頭の＊を削除（拒否の連絡先を示す定型文なので行は残す）", "＊"]]
        elif kinds == {"red"}:
            reason = "赤字（記載上の注意）"
        elif "black" not in kinds:
            reason = "青字（例文）"
        elif t.startswith(("＊", "*")) or any(k in seg_text(s) for s in segs if s[0] == "black" for k in BLACK_EXAMPLE):
            reason = "黒字の例文（＊・○○）"
        if reason:
            if current is not None and current not in style_ref:
                style_ref[current] = copy.deepcopy(b._p)
            stats["removed"].append({"page": page, "reason": reason, "text": t})
            remove(b)
            continue
        log = star_log + apply_inline(b)
        if b.text != before:
            stats["partial"].append({"page": page, "before": before.strip(), "after": b.text.strip(), "changes": log})
    return headings, style_ref


def add_placeholders(doc, headings, style_ref, stats):
    """文書の順に見出しと表を見て、本文が空の見出しの下に（本文を記載）を置く。"""
    bl = blocks(doc)
    for i, b in enumerate(bl):
        if isinstance(b, Table):
            clean_table(b, headings.get(b._tbl, (None, "table"))[0], stats)
            continue
        if b._p not in headings:
            continue
        page, kind = headings[b._p]
        t = b.text.strip()
        row = {"page": page, "heading": t, "kind": kind}
        stats["headings"].append(row)
        if kind == "title":
            row["status"] = "表題（下は黒字の定型文）"
            continue
        if kind == "h1" and t.split("]", 1)[1].strip():
            row["status"] = "同じ行に記載欄: " + t.split("]", 1)[1].strip()
            continue
        nxt = None
        for nb in bl[i + 1:]:
            if isinstance(nb, Table):
                nxt = nb
                break
            if nb.text.strip():
                nxt = nb
                break
            if has_page_break(nb):
                break
        if isinstance(nxt, Table):
            row["status"] = "下に表"
            continue
        if nxt is not None:
            nk = heading_kind(nxt.text)
            if kind == "h1" and nk == "h2":
                row["status"] = "下の小見出しに続く"
                continue
            if nk is None:
                row["status"] = "本文（黒字の定型文）あり: " + nxt.text.strip()[:40]
                continue
        ref = style_ref.get(b._p)
        ref = Paragraph(ref if ref is not None else copy.deepcopy(b._p), None)
        for m in insert_after(b, [PLACEHOLDER], ref_p=ref, size=BODY_PT, parent=doc._body):
            clear_mark_color(m)
        row["status"] = PLACEHOLDER + " を挿入"
        stats["placeholders"].append({"page": page, "where": t})


def clean_table(tbl, page, stats):
    """表の見出し行は残して黒に、「例）」の行は消し、最初の空の行に（本文を記載）。"""
    header = " | ".join(c.text.strip() for c in distinct_cells(tbl.rows[0]))
    for row in list(tbl.rows):
        first = distinct_cells(row)[0].text.strip()
        if first.startswith(("例）", "例)", "（例", "(例")):
            stats["table_rows_removed"].append({"page": page, "table": header,
                                                "row": " | ".join(c.text.strip() for c in distinct_cells(row))})
            row._tr.getparent().remove(row._tr)
    for row in tbl.rows:
        for c in distinct_cells(row):
            for p in c.paragraphs:
                for r in list(p.runs):
                    if cls_of(r) == "red":
                        remove(r)
                    elif cls_of(r) == "blue":
                        clear_marks(r)          # 見出し行（研究機関名 など）は表の見出しとして残す
                clear_mark_color(p)
    status = "見出し行を残した（空の行なし）"
    for row in list(tbl.rows)[1:]:
        cells = distinct_cells(row)
        if all(not c.text.strip() for c in cells):
            r = cells[0].paragraphs[0].add_run(PLACEHOLDER)
            r.font.size = Pt(BODY_PT)
            stats["placeholders"].append({"page": page, "where": "表: " + header})
            status = "見出し行を残し、空の行の1列目に" + PLACEHOLDER
            break
    stats["headings"].append({"page": page, "heading": "（表）" + header, "kind": "table", "status": status})


def clean_headers(doc, stats):
    for sec in doc.sections:
        parts = [("header", sec.header), ("first_page_header", sec.first_page_header),
                 ("even_page_header", sec.even_page_header), ("footer", sec.footer),
                 ("first_page_footer", sec.first_page_footer), ("even_page_footer", sec.even_page_footer)]
        for name, part in parts:
            if part.is_linked_to_previous:        # 定義のない部品に触ると空の部品が作られるので触らない
                continue
            for p in part.paragraphs:
                before = p.text
                apply_inline(p)
                for r in p.runs:
                    if re.search(r"[〇○◯]", r.text):
                        r.text = re.sub(r"[〇○◯]+", SLOT, r.text)
                    if cls_of(r) != "black":
                        clear_marks(r)
                if p.text != before:
                    stats["header"].append({"part": name, "before": before, "after": p.text})
        first = sec.first_page_header
        if sec.different_first_page_header_footer and not first.is_linked_to_previous \
                and not any(p.text.strip() for p in first.paragraphs):
            sec.different_first_page_header_footer = False
            stats["header"].append({"part": "section", "before": "先頭ページのみ別指定（空ヘッダー）",
                                    "after": "指定を外し、全ページに作成日・版数のヘッダー"})


# ------------------------------------------------------------------ 確認
def verify(path):
    doc = Document(str(path))
    problems, n_ph, n_slot = [], 0, 0

    def look(where, p):
        nonlocal n_ph, n_slot
        t = p.text
        n_ph += t.count(PLACEHOLDER)
        n_slot += t.count(SLOT)
        for r in p.runs:
            if r.text.strip() and cls_of(r) != "black":
                problems.append(f"{where}: 色が残る {r.text[:20]!r}")
        ppr = p._p.pPr
        rpr = None if ppr is None else ppr.find(qn("w:rPr"))
        col = None if rpr is None else rpr.find(qn("w:color"))
        if col is not None and (col.get(qn("w:val")) or "").upper() in RED | GUIDE_BLUE:
            problems.append(f"{where}: 段落記号に色 {t[:20]!r}")
        for k in TRACE:
            if k in t.replace(SLOT, "") and not t.strip().startswith(("○", "〇")):
                problems.append(f"{where}: 痕跡「{k}」 {t.strip()[:40]!r}")

    for i, b in enumerate(blocks(doc)):
        if isinstance(b, Paragraph):
            look(f"#{i}", b)
        else:
            for row in b.rows:
                for c in distinct_cells(row):
                    for p in c.paragraphs:
                        look(f"表#{i}", p)
    for sec in doc.sections:
        for name in ("header", "first_page_header", "footer"):
            part = getattr(sec, name)
            if not part.is_linked_to_previous:
                for p in part.paragraphs:
                    look(name, p)
    return problems, n_ph, n_slot, doc.sections[0].different_first_page_header_footer


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, help="北大の HT3 雛形（HT3_情報公開用文書ひな形_202504.docx）")
    ap.add_argument("--variant", choices=("both", "single", "multi"), default="both")
    ap.add_argument("--out")
    ap.add_argument("--report", help="処理の記録（JSON）の保存先")
    a = ap.parse_args()
    global TPL
    TPL = Path(a.template)
    suffix = {"both": "", "single": "_単独", "multi": "_多機関"}[a.variant]
    out = Path(a.out).resolve() if a.out else Path.cwd() / f"HT3_情報公開用文書_下書き{suffix}.docx"
    if out == TPL.resolve() or out.parent == TPL.resolve().parent:
        sys.exit(f"ここには書きません: {out}")

    stats = {"template": str(TPL), "variant": a.variant, "removed": [], "partial": [], "placeholders": [],
             "table_rows_removed": [], "headings": [], "header": []}
    doc = Document(str(TPL))
    cut_variant(doc, a.variant, stats)
    headings, style_ref = purge(doc, stats)
    add_placeholders(doc, headings, style_ref, stats)     # 表の片付けもここで（文書の順）
    clean_headers(doc, stats)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    problems, n_ph, n_slot, title_pg = verify(out)
    stats.update({"output": str(out), "placeholder_count": n_ph, "slot_count": n_slot, "problems": problems})

    # ---- 報告
    print(f"保存: {out}")
    cov = stats["cover"]
    print(f"\n冒頭の使用上の注意ページ: {cov['blocks']}段落（文字あり {cov['nonempty']}・空行 {cov['blocks'] - cov['nonempty']}）")
    if "other_variant" in stats:
        ov = stats["other_variant"]
        print(f"使わない方のページ: {ov['blocks']}段落・表（文字あり {ov['nonempty']}）")
    by = {}
    for r in stats["removed"]:
        by.setdefault(r["reason"], []).append(r)
    print(f"\n本文で段落ごと消したもの: {len(stats['removed'])}段落")
    for k, v in by.items():
        print(f"  {k}: {len(v)}")
        for r in v:
            print(f"     {r['page']} {r['text'][:50]}")
    print(f"\n行は残して一部だけ消した/置き換えたもの: {len(stats['partial'])}段落")
    for r in stats["partial"]:
        print(f"     {r['page']} {r['before'][:34]}  →  {r['after'][:40]}")
    print(f"\n表の例示行を消した: {len(stats['table_rows_removed'])}行")
    for r in stats["table_rows_removed"]:
        print(f"     {r['page']} {r['row']}")
    print("\nヘッダー:")
    for r in stats["header"]:
        print(f"     {r['part']}: {r['before']} → {r['after']}")
    print("\n残した見出し:")
    for r in stats["headings"]:
        print(f"     {r['page']} {r['heading'][:40]:<40} … {r['status']}")
    print(f"\n（本文を記載）: {n_ph}か所　【　　】: {n_slot}か所　先頭ページ別ヘッダー: {title_pg}")
    print("確認:", "問題なし" if not problems else f"{len(problems)}件")
    for x in problems:
        print("   !!", x)
    if a.report:
        Path(a.report).write_text(json.dumps(stats, ensure_ascii=False, indent=1), encoding="utf-8")
        print("記録:", a.report)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
