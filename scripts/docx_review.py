# -*- coding: utf-8 -*-
"""Word の「校閲（変更履歴）」と「コメント」で書類をやりとりする道具。

ターミナルを使わない人とも、Word だけで直しを回せるようにする:
  1. 相手が Word で気になる所にコメントを付けて渡す
  2. comments で読み、直し方と返事を plan.json に書く
  3. apply で、直しを**変更履歴**として書き込み、各コメントに**返信**する（解決済みの印も付けられる）
  4. 相手は Word で「承諾／元に戻す」を選び、返信を読む。また気になればコメントを付けて返す
元のファイルは上書きしない（別名で保存する）。人が Word で直した版を正本として扱うため。

使い方:
  python docx_review.py comments 計画書.docx [--json comments.json]
  python docx_review.py apply 計画書.docx plan.json [--out 計画書_校閲.docx]
  python docx_review.py revisions 計画書.docx          # 残っている変更履歴の一覧
  python docx_review.py accept-all 計画書.docx --out 計画書_清書.docx   # 全部承諾した版

plan.json の形（どれも省略できる）:
{
  "author": "Claude", "initials": "C",
  "edits": [
    {"para": "^6．", "find": "予測率", "replace": "変化率", "comment": "13章の解析と量をそろえました"},
    {"para": "症例数", "find": "約100例", "insert": "（2021〜2023年の実績）"},
    {"insert_after": "^【設定根拠】$", "text": "1段落目\\n2段落目", "style_from": "next"},
    {"delete_paragraph": "^＊例：", "comment": "ひな形の例文なので削除"}
  ],
  "replies": [
    {"comment_id": 3, "text": "ご指摘のとおり直しました（変更履歴をご確認ください）。", "resolve": true}
  ],
  "new_comments": [
    {"para": "^13．", "anchor": "目標症例数", "text": "根拠の数値（過去3年の実績）をご確認ください。"}
  ]
}
- para / insert_after / delete_paragraph は段落の文字に対する正規表現（表のセルの段落も探す）。1つに決まらないと止まる（"all": true で全部）。
- find は段落の中の文字そのまま（正規表現にしたいときは "regex": true）。段落の中で1か所に決まらないと止まる。
- 直しは段落の中だけ（段落をまたぐ置き換えはしない）。複雑な段落（既存の変更履歴の中、タブ・改行をまたぐ）は止まるので、手で直すか段落ごと差し替える。
"""
import argparse
import copy
import datetime as dt
import io
import json
import random
import re
import sys
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
XML = "http://www.w3.org/XML/1998/namespace"
RT_CEX = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
CT_CEX = "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"
NS = {"w": W, "w14": W14, "w15": W15}


def xp(e, expr):
    """lxml の xpath（python-docx の要素は名前空間の渡し方が違うので、元の関数を呼ぶ）。"""
    return etree._Element.xpath(e, expr, namespaces=NS)


def w(tag):
    return f"{{{W}}}{tag}"


def el(tag, **attrs):
    e = etree.Element(w(tag))
    for k, v in attrs.items():
        e.set(w(k), str(v))
    return e


def now():
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


# ── 読む ────────────────────────────────────────────────

def has_comments(doc):
    return any(r.reltype == RT.COMMENTS for r in doc.part.rels.values())


def comments_root(doc):
    """comments.xml の根（無ければ作る）。"""
    return doc.part._comments_part.element


def cex_part(doc, create=False):
    for r in doc.part.rels.values():
        if r.reltype == RT_CEX:
            return r.target_part
    if not create:
        return None
    xml = (f'<w15:commentsEx xmlns:w15="{W15}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           f'mc:Ignorable="w15"/>').encode("utf-8")
    part = Part(PackURI("/word/commentsExtended.xml"), CT_CEX, xml, doc.part.package)
    doc.part.relate_to(part, RT_CEX)
    return part


def cex_root(part):
    if hasattr(part, "_element"):
        return part._element
    return etree.fromstring(part.blob)


def cex_save(part, root):
    if hasattr(part, "_element"):
        return
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def text_of(e):
    return "".join(t.text or "" for t in e.iter(w("t")))


def para_paraid(p):
    return p.get(f"{{{W14}}}paraId")


def comment_info(doc):
    """コメントの一覧（スレッドと解決済みの印つき）。"""
    if not has_comments(doc):
        return []
    root = comments_root(doc)
    items = {}
    for c in root.iter(w("comment")):
        cid = c.get(w("id"))
        ps = c.findall(w("p"))
        items[cid] = {"id": int(cid), "author": c.get(w("author")), "date": c.get(w("date")),
                      "text": "\n".join(text_of(p) for p in ps).strip(),
                      "para_id": para_paraid(ps[-1]) if ps else None,
                      "anchor": "", "paragraph": "", "parent_id": None, "done": False}
    # 本文のどこに付いているか
    body = doc.element.body
    active, start_para = set(), {}
    for e in body.iter():
        if e.tag == w("commentRangeStart"):
            cid = e.get(w("id"))
            active.add(cid)
            p = next((a for a in e.iterancestors(w("p"))), None)
            if cid in items and p is not None:
                items[cid]["paragraph"] = "".join(t.text or "" for t in p.iter(w("t"), w("delText")))
        elif e.tag == w("commentRangeEnd"):
            active.discard(e.get(w("id")))
        elif e.tag in (w("t"), w("delText")):
            deleted = e.tag == w("delText")
            for cid in active:
                if cid in items:
                    items[cid]["anchor"] += f"〔削除: {e.text or ''}〕" if deleted else (e.text or "")
    # スレッドと解決済み
    part = cex_part(doc)
    if part is not None:
        by_para = {v["para_id"]: k for k, v in items.items() if v["para_id"]}
        for ex in cex_root(part).iter(f"{{{W15}}}commentEx"):
            cid = by_para.get(ex.get(f"{{{W15}}}paraId"))
            if cid is None:
                continue
            items[cid]["done"] = ex.get(f"{{{W15}}}done") == "1"
            parent = ex.get(f"{{{W15}}}paraIdParent")
            if parent and parent in by_para:
                items[cid]["parent_id"] = int(by_para[parent])
    return sorted(items.values(), key=lambda x: x["id"])


def cmd_comments(a):
    doc = Document(a.docx)
    items = comment_info(doc)
    if not items:
        print("コメントはありません。")
    roots = [c for c in items if c["parent_id"] is None]
    for c in roots:
        mark = "［解決済み］" if c["done"] else ""
        print(f"#{c['id']} {c['author']} {c['date'] or ''} {mark}")
        print(f"   場所: 「{c['anchor'][:60]}」 （段落: {c['paragraph'][:50]}）")
        print(f"   {c['text']}")
        for r in [x for x in items if x["parent_id"] == c["id"]]:
            print(f"   └ #{r['id']} {r['author']}: {r['text']}")
    if a.json:
        Path(a.json).write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"→ {a.json}")


# ── 書く（変更履歴とコメント） ───────────────────────────────

class Ids:
    """変更履歴とコメントの番号を、文書の中で重ならないように振る。"""

    def __init__(self, doc):
        nums = [int(v) for v in xp(doc.element.body, "//@w:id") if str(v).lstrip("-").isdigit()]
        self.rev = max(nums + [0]) + 1000
        self.com = 0
        self.para = set(xp(doc.element.body, "//@w14:paraId"))
        if has_comments(doc):
            root = comments_root(doc)
            cids = [int(c.get(w("id"))) for c in root.iter(w("comment"))]
            self.com = max(cids + [-1]) + 1
            self.para |= set(xp(root, "//@w14:paraId"))

    def next_rev(self):
        self.rev += 1
        return self.rev

    def next_com(self):
        n = self.com
        self.com += 1
        return n

    def new_para(self):
        while True:
            v = f"{random.randint(1, 0x7FFFFFFF):08X}"
            if v not in self.para:
                self.para.add(v)
                return v


def body_paras(doc):
    return list(doc.element.body.iter(w("p")))


def pick_paras(doc, rx, allow_many):
    hits = [p for p in body_paras(doc) if re.search(rx, text_of(p))]
    if not hits:
        raise SystemExit(f"段落が見つからない: {rx}")
    if len(hits) > 1 and not allow_many:
        raise SystemExit(f"段落が {len(hits)} 個ある（\"all\": true で全部、または正規表現を絞る）: {rx}")
    return hits


def direct_runs(p):
    return [r for r in p if r.tag == w("r")]


def run_text(r):
    return "".join(t.text or "" for t in r.findall(w("t")))


def _plain(r):
    """w:t だけ（と rPr）の単純なラン。タブ・改行・図などを含むランは割らない。"""
    return all(c.tag in (w("rPr"), w("t"), w("lastRenderedPageBreak")) for c in r)


def _set_text(r, s):
    ts = r.findall(w("t"))
    for t in ts[1:]:
        r.remove(t)
    t = ts[0] if ts else etree.SubElement(r, w("t"))
    t.text = s
    t.set(f"{{{XML}}}space", "preserve")


def split_at(p, pos):
    """段落 p の文字位置 pos でランを割り、pos から始まるランを返す（pos が末尾なら None）。"""
    i = 0
    for r in direct_runs(p):
        s = run_text(r)
        if i == pos:
            return r
        if i < pos < i + len(s):
            if not _plain(r):
                raise SystemExit("割れないラン（タブ・改行・図など）の中を指している。段落ごと差し替えてほしい")
            r2 = copy.deepcopy(r)
            _set_text(r, s[: pos - i])
            _set_text(r2, s[pos - i:])
            r.addnext(r2)
            return r2
        i += len(s)
    return None


def find_span(p, needle, regex):
    full = "".join(run_text(r) for r in direct_runs(p))
    ms = list(re.finditer(needle if regex else re.escape(needle), full))
    if not ms:
        if needle in text_of(p):
            raise SystemExit(f"「{needle}」は既存の変更履歴やリンクの中にある。手で直してほしい")
        raise SystemExit(f"段落の中に見つからない: 「{needle}」（段落: {text_of(p)[:40]}）")
    if len(ms) > 1:
        raise SystemExit(f"段落の中に {len(ms)} か所ある（前後の語を足して1か所にする）: 「{needle}」")
    return ms[0].start(), ms[0].end()


def runs_between(p, a, b):
    first = split_at(p, a)
    split_at(p, b)
    out, i = [], 0
    for r in direct_runs(p):
        s = run_text(r)
        if a <= i < b:
            out.append(r)
        i += len(s)
    return first, out


def mark(tag, ids, author, date):
    return el(tag, id=ids.next_rev(), author=author, date=date)


def to_deleted(r):
    for t in r.findall(w("t")):
        t.tag = w("delText")
        t.set(f"{{{XML}}}space", "preserve")


def new_run(text, rpr=None):
    r = el("r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = etree.SubElement(r, w("t"))
    t.text = text
    t.set(f"{{{XML}}}space", "preserve")
    return r


def tracked_replace(p, find, replace, insert, regex, ids, author, date):
    """段落の中の find を、変更履歴つきで replace に置き換える（insert なら後ろに足す）。
    返り値は (先頭の要素, 末尾の要素)＝コメントを付ける範囲。"""
    a, b = find_span(p, find, regex)
    first, rs = runs_between(p, a, b)
    rpr = rs[0].find(w("rPr")) if rs else None
    if insert is not None:
        ins = mark("ins", ids, author, date)
        ins.append(new_run(insert, rpr))
        rs[-1].addnext(ins)
        return ins, ins
    d = mark("del", ids, author, date)
    rs[0].addprevious(d)
    for r in rs:
        to_deleted(r)
        d.append(r)
    last = d
    if replace:
        ins = mark("ins", ids, author, date)
        ins.append(new_run(replace, rpr))
        d.addnext(ins)
        last = ins
    return d, last


def ppr_mark(p, tag, ids, author, date):
    ppr = p.find(w("pPr"))
    if ppr is None:
        ppr = el("pPr")
        p.insert(0, ppr)
    rpr = ppr.find(w("rPr"))
    if rpr is None:
        rpr = etree.SubElement(ppr, w("rPr"))
    rpr.append(mark(tag, ids, author, date))


def tracked_insert_paras(p, text, style_from, ids, author, date):
    src = p
    if style_from == "next":
        nxt = p.getnext()
        src = nxt if nxt is not None and nxt.tag == w("p") else p
    base_ppr = src.find(w("pPr")) if style_from != "none" else None
    base_rpr = next((r.find(w("rPr")) for r in direct_runs(src) if r.find(w("rPr")) is not None), None)
    anchor, made = p, []
    for line in text.split("\n"):
        np_ = el("p")
        if base_ppr is not None:
            pp = copy.deepcopy(base_ppr)
            for old in pp.findall(w("rPr")):
                pp.remove(old)
            for sect in pp.findall(w("sectPr")):
                pp.remove(sect)
            np_.append(pp)
        ppr_mark(np_, "ins", ids, author, date)
        ins = mark("ins", ids, author, date)
        ins.append(new_run(line, base_rpr))
        np_.append(ins)
        anchor.addnext(np_)
        anchor = np_
        made.append(np_)
    return made


def tracked_delete_para(p, ids, author, date):
    rs = direct_runs(p)
    if rs:
        d = mark("del", ids, author, date)
        rs[0].addprevious(d)
        for r in rs:
            to_deleted(r)
            d.append(r)
    ppr_mark(p, "del", ids, author, date)


def _comment_xml(cid, author, initials, date, text, para_id):
    c = el("comment", id=cid, author=author, date=date, initials=initials)
    lines = text.split("\n") or [""]
    for k, line in enumerate(lines):
        p = el("p")
        if k == len(lines) - 1:
            p.set(f"{{{W14}}}paraId", para_id)
            p.set(f"{{{W14}}}textId", "77777777")
        ppr = el("pPr")
        st = el("pStyle", val="CommentText")
        ppr.append(st)
        p.append(ppr)
        if k == 0:
            r0 = el("r")
            rpr = el("rPr")
            rpr.append(el("rStyle", val="CommentReference"))
            r0.append(rpr)
            r0.append(el("annotationRef"))
            p.append(r0)
        p.append(new_run(line))
        c.append(p)
    return c


def _ref_run(cid):
    r = el("r")
    rpr = el("rPr")
    rpr.append(el("rStyle", val="CommentReference"))
    r.append(rpr)
    r.append(el("commentReference", id=cid))
    return r


def add_comment(doc, ids, start_el, end_el, text, author, initials, date, parent=None):
    """start_el の前から end_el の後ろまでにコメントを付ける。parent があれば、そのコメントへの返信にする。"""
    root = comments_root(doc)
    cid = ids.next_com()
    pid = ids.new_para()
    root.append(_comment_xml(cid, author, initials, date, text, pid))
    if parent is None:
        start_el.addprevious(el("commentRangeStart", id=cid))
        end_el.addnext(el("commentRangeEnd", id=cid))
        end = end_el.getnext()
        end.addnext(_ref_run(cid))
    else:
        body = doc.element.body
        ps = xp(body, f'.//w:commentRangeStart[@w:id="{parent["id"]}"]')
        pe = xp(body, f'.//w:commentRangeEnd[@w:id="{parent["id"]}"]')
        pr = xp(body, f'.//w:r[w:commentReference[@w:id="{parent["id"]}"]]')
        if not (ps and pe and pr):
            raise SystemExit(f"コメント #{parent['id']} の位置が本文に見つからない")
        ps[0].addnext(el("commentRangeStart", id=cid))
        pe[0].addnext(el("commentRangeEnd", id=cid))
        pr[0].addnext(_ref_run(cid))
    return cid, pid


def ensure_para_id(doc, ids, comment_id):
    root = comments_root(doc)
    c = xp(root, f'.//w:comment[@w:id="{comment_id}"]')
    if not c:
        raise SystemExit(f"コメント #{comment_id} が無い")
    ps = c[0].findall(w("p"))
    if not ps:
        raise SystemExit(f"コメント #{comment_id} に段落が無い")
    pid = para_paraid(ps[-1])
    if not pid:
        pid = ids.new_para()
        ps[-1].set(f"{{{W14}}}paraId", pid)
        ps[-1].set(f"{{{W14}}}textId", "77777777")
    return pid


def cex_set(doc, para_id, parent_para=None, done=None):
    part = cex_part(doc, create=True)
    root = cex_root(part)
    ex = None
    for e in root.iter(f"{{{W15}}}commentEx"):
        if e.get(f"{{{W15}}}paraId") == para_id:
            ex = e
    if ex is None:
        ex = etree.SubElement(root, f"{{{W15}}}commentEx")
        ex.set(f"{{{W15}}}paraId", para_id)
        ex.set(f"{{{W15}}}done", "0")
    if parent_para:
        ex.set(f"{{{W15}}}paraIdParent", parent_para)
    if done is not None:
        ex.set(f"{{{W15}}}done", "1" if done else "0")
    cex_save(part, root)


def cmd_apply(a):
    src = Path(a.docx)
    out = Path(a.out) if a.out else src.with_name(f"{src.stem}_校閲_{dt.datetime.now():%Y%m%d-%H%M}.docx")
    if out.resolve() == src.resolve():
        raise SystemExit("元のファイルには書かない（--out に別名を）")
    plan = json.loads(Path(a.plan).read_text(encoding="utf-8"))
    author, initials = plan.get("author", "Claude"), plan.get("initials", "C")
    date = now()
    doc = Document(str(src))
    ids = Ids(doc)
    log = []

    for e in plan.get("edits", []):
        many = bool(e.get("all"))
        if "delete_paragraph" in e:
            for p in pick_paras(doc, e["delete_paragraph"], many):
                txt = text_of(p)
                tracked_delete_para(p, ids, author, date)
                if e.get("comment"):
                    rs = p.findall(w("del"))
                    if rs:
                        add_comment(doc, ids, rs[0], rs[-1], e["comment"], author, initials, date)
                log.append(f"段落を削除: {txt[:40]}")
        elif "insert_after" in e:
            for p in pick_paras(doc, e["insert_after"], many):
                made = tracked_insert_paras(p, e["text"], e.get("style_from", "next"), ids, author, date)
                if e.get("comment"):
                    first = made[0].find(w("ins"))
                    last = made[-1].find(w("ins"))
                    add_comment(doc, ids, first, last, e["comment"], author, initials, date)
                log.append(f"段落を追加（{text_of(p)[:20]} の後）: {e['text'][:40]}")
        elif "find" in e:
            paras = pick_paras(doc, e["para"], many) if e.get("para") else None
            if paras is None:
                paras = [p for p in body_paras(doc) if e["find"] in "".join(run_text(r) for r in direct_runs(p))]
                if len(paras) != 1 and not many:
                    raise SystemExit(f"「{e['find']}」を含む段落が {len(paras)} 個（\"para\" で段落を絞る）")
            for p in paras:
                s, t = tracked_replace(p, e["find"], e.get("replace", ""), e.get("insert"), bool(e.get("regex")),
                                       ids, author, date)
                if e.get("comment"):
                    add_comment(doc, ids, s, t, e["comment"], author, initials, date)
                what = f"「{e['find']}」→「{e.get('replace', '')}」" if e.get("insert") is None else f"「{e['find']}」の後に「{e['insert']}」"
                log.append(f"直し: {what}")
        else:
            raise SystemExit(f"edits の形が分からない: {e}")

    for nc in plan.get("new_comments", []):
        p = pick_paras(doc, nc["para"], False)[0]
        if nc.get("anchor"):
            first, rs = runs_between(p, *find_span(p, nc["anchor"], bool(nc.get("regex"))))
            start, end = rs[0], rs[-1]
        else:
            rs = direct_runs(p)
            if not rs:
                raise SystemExit(f"空の段落にはコメントを付けられない: {nc['para']}")
            start, end = rs[0], rs[-1]
        add_comment(doc, ids, start, end, nc["text"], author, initials, date)
        log.append(f"コメント: {nc['text'][:40]}")

    info = {c["id"]: c for c in comment_info(doc)}
    for rp in plan.get("replies", []):
        cid = int(rp["comment_id"])
        if cid not in info:
            raise SystemExit(f"コメント #{cid} が無い（comments で番号を確かめる）")
        parent_para = ensure_para_id(doc, ids, cid)
        new_id, new_para = add_comment(doc, ids, None, None, rp["text"], author, initials, date, parent=info[cid])
        cex_set(doc, parent_para, done=bool(rp.get("resolve")) if "resolve" in rp else None)
        cex_set(doc, new_para, parent_para=parent_para)
        log.append(f"返信 #{cid}{'（解決済み）' if rp.get('resolve') else ''}: {rp['text'][:40]}")

    doc.save(str(out))
    print("\n".join(log) or "（何もしていない）")
    print(f"→ {out}")


# ── 確かめる・清書 ───────────────────────────────────────────

def cmd_revisions(a):
    doc = Document(a.docx)
    n = 0
    for e in doc.element.body.iter(w("ins"), w("del")):
        txt = "".join((t.text or "") for t in e.iter(w("t"), w("delText")))
        if not txt and e.getparent() is not None and e.getparent().tag == w("rPr"):
            kind = "段落記号の" + ("追加" if e.tag == w("ins") else "削除")
        else:
            kind = "追加" if e.tag == w("ins") else "削除"
        print(f"{kind:8} {e.get(w('author'))} {e.get(w('date'))}: {txt[:60]}")
        n += 1
    print(f"変更履歴 {n} か所")


def cmd_accept_all(a):
    src = Path(a.docx)
    out = Path(a.out)
    if out.resolve() == src.resolve():
        raise SystemExit("元のファイルには書かない（--out に別名を）")
    doc = Document(str(src))
    body = doc.element.body
    # 削除された段落（段落記号ごと消えたもの）を先に消す
    for p in list(body.iter(w("p"))):
        ppr = p.find(w("pPr"))
        rpr = ppr.find(w("rPr")) if ppr is not None else None
        if rpr is not None and rpr.find(w("del")) is not None and not "".join(t.text or "" for t in p.iter(w("t"))
                                                                               if not any(x.tag == w("del") for x in t.iterancestors())):
            p.getparent().remove(p)
    for tag in ("del", "moveFrom"):
        for e in list(body.iter(w(tag))):
            if e.getparent() is not None:
                e.getparent().remove(e)
    for tag in ("ins", "moveTo"):
        for e in list(body.iter(w(tag))):
            parent = e.getparent()
            if parent is None:
                continue
            if parent.tag == w("rPr"):
                parent.remove(e)
                continue
            i = parent.index(e)
            for ch in list(e):
                parent.insert(i, ch)
                i += 1
            parent.remove(e)
    for tag in ("rPrChange", "pPrChange", "sectPrChange", "tblPrChange", "trPrChange", "tcPrChange",
                "moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd"):
        for e in list(body.iter(w(tag))):
            e.getparent().remove(e)
    doc.save(str(out))
    print(f"全部承諾した版 → {out}")


def main():
    ap = argparse.ArgumentParser(description="Word の変更履歴とコメントでやりとりする")
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("comments")
    c.add_argument("docx")
    c.add_argument("--json")
    c.set_defaults(fn=cmd_comments)
    p = sub.add_parser("apply")
    p.add_argument("docx")
    p.add_argument("plan")
    p.add_argument("--out")
    p.set_defaults(fn=cmd_apply)
    r = sub.add_parser("revisions")
    r.add_argument("docx")
    r.set_defaults(fn=cmd_revisions)
    x = sub.add_parser("accept-all")
    x.add_argument("docx")
    x.add_argument("--out", required=True)
    x.set_defaults(fn=cmd_accept_all)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
