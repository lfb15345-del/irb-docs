# -*- coding: utf-8 -*-
"""Wordひな形（または記入済みの書類）の中身を、色・見出し・表・テキストボックス・ヘッダー/フッターつきで書き出す。

転記スクリプトを書く前に必ず一度流す。どの行が赤（注意）・青（例文）・黒（定型文）か、
見出しの正確な文字と深さ、表の結合セル、表紙の枠（テキストボックス）、先頭ページ別指定がここで分かる。

    python dump_template.py ひな形.docx [--out 結果.txt] [--grep 語] [--runs] [--full]

  --runs  1行の中で色が混ざる行について、色ごとの文字のかたまりを出す（mode="inline" が要るか判断する）
  --full  段落を切らずに全文を出す（記入済みの書類の点検用）

.doc は読めないので、先に docx_to_pdf.ps1 -ToDocx で .docx にする。
"""
import argparse
import io
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.table import Table  # noqa: E402

from docx_tools import blocks, chrome_parts, distinct_cells, heading_level, run_class, segments  # noqa: E402

MARK = {"red": "赤", "blue": "青", "black": "黒"}


def color_tag(p):
    cols = Counter(run_class(r, p) for r in p.runs if r.text.strip())
    if not cols:
        return "空"
    return "+".join(MARK[c] for c, _ in cols.most_common())


def run_view(p):
    parts = []
    for c, rs in segments(p):
        text = "".join(r.text for r in rs)
        if text.strip():
            parts.append(f"{MARK[c]}«{text}»")
    return " ".join(parts)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--out")
    ap.add_argument("--grep", help="この語を含む行だけ出す")
    ap.add_argument("--runs", action="store_true", help="色が混ざる行の中身を色ごとに出す")
    ap.add_argument("--full", action="store_true", help="段落を切らずに出す")
    a = ap.parse_args()
    width = 100000 if a.full else 110
    doc = Document(a.docx)
    lines = []
    for si, sec in enumerate(doc.sections):
        lines.append(f"[セクション{si}] 先頭ページ別指定={sec.different_first_page_header_footer}"
                     f" 奇数偶数別={doc.settings.odd_and_even_pages_header_footer}")
    for i, b in enumerate(blocks(doc)):
        if isinstance(b, Table):
            lines.append(f"{i:04d} [表 {len(b.rows)}行]")
            for ri, row in enumerate(b.rows):
                cells = distinct_cells(row)
                texts = [c.text.strip().replace("\n", "/")[:width // 2] for c in cells]
                kinds = {run_class(r, p) for c in cells for p in c.paragraphs for r in p.runs if r.text.strip()}
                mark = " *色*" if kinds - {"black"} else ""
                lines.append(f"      r{ri}: " + " | ".join(texts) + mark)
            continue
        t = b.text.strip()
        if not t:
            if b._p.findall(".//" + qn("w:sectPr")):
                lines.append(f"{i:04d} [セクション区切り]")
            elif any(br.get(qn("w:type")) == "page" for br in b._p.iter(qn("w:br"))):
                lines.append(f"{i:04d} [改ページ]")
            continue
        lv = heading_level(t)
        head = f"H{lv:g}" if lv else "  "
        style = b.style.name if b.style is not None else ""
        tag = color_tag(b)
        lines.append(f"{i:04d} {head:<4} [{tag}] ({style}) {t[:width]}")
        if a.runs and "+" in tag:
            lines.append(f"         └ {run_view(b)[:width * 2]}")
    tbs = doc.element.body.findall(".//" + qn("w:txbxContent"))
    if tbs:
        lines.append("--- テキストボックス")
        for tb in tbs:
            txt = " / ".join("".join(n.text or "" for n in p.findall(".//" + qn("w:t"))) for p in tb.findall(qn("w:p")))
            lines.append("   " + txt[:width])
    lines.append("--- ヘッダー/フッター（定義のあるものだけ）")
    for name, part in chrome_parts(doc):
        for p in part.paragraphs:
            if p.text.strip():
                lines.append(f"   {name}: [{color_tag(p)}] {p.text.strip()[:width]}")
                if a.runs and "+" in color_tag(p):
                    lines.append(f"         └ {run_view(p)}")
    if a.grep:
        lines = [x for x in lines if a.grep in x]
    text = "\n".join(lines)
    if a.out:
        Path(a.out).write_text(text, encoding="utf-8")
        print(f"saved {a.out}（{len(lines)} 行）")
    else:
        print(text)


if __name__ == "__main__":
    main()
